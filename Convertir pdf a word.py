import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def establecer_fondo_celda(celda, color_hex):
    """Permite pintar el fondo de una celda de Word con color hexadecimal (ej. '333333')."""
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def extraer_y_construir():
    ruta_base = "/mnt/c/Users/isamu/OneDrive/Desktop/Trabajos Álvaro Flores/Cosas"
    pdf_path = os.path.join(ruta_base, "C.V. Álvaro Flores.pdf")
    docx_path = os.path.join(ruta_base, "C.V. Álvaro Flores_pro.docx")

    print(f"Leyendo coordenadas geométricas de: {pdf_path}")
    doc_pdf = fitz.open(pdf_path)
    pagina = doc_pdf[0]

    # Extraer bloques de texto con sus coordenadas (x0, y0, x1, y1, texto, block_no, block_type)
    bloques = pagina.get_text("blocks")

    # Inicializar documento Word
    doc_word = Document()
    
    # Configurar márgenes estrechos para aprovechar la hoja igual que el PDF
    for seccion in doc_word.sections:
        seccion.top_margin = Inches(0.5)
        seccion.bottom_margin = Inches(0.5)
        seccion.left_margin = Inches(0.5)
        seccion.right_margin = Inches(0.5)

    # Creamos una tabla principal de 2 columnas para simular el layout del C.V.
    # Columna izquierda (principal), Columna derecha (lateral o viceversa según diseño)
    tabla = doc_word.add_table(rows=1, cols=2)
    tabla.autofit = False
    
    # Ajustar anchos aproximados de las columnas del currículum (ej. 60% y 40%)
    tabla.columns[0].width = Inches(4.0)
    tabla.columns[1].width = Inches(2.5)

    celda_izq = tabla.cell(0, 0)
    celda_der = tabla.cell(0, 1)

    # Ordenar bloques verticalmente por su posición 'y' (y0)
    bloques_ordenados = sorted(bloques, key=lambda b: b[1])

    for b in bloques_ordenados:
        x0, y0, x1, y1, texto, block_no, block_type = b
        texto_limpio = texto.strip()
        
        if not texto_limpio:
            continue

        # Separar contenido según la coordenada X horizontal para meterlo en columna izquierda o derecha
        # Si x0 es menor a la mitad de la página, va a la columna izquierda, sino a la derecha
        p = celda_izq.add_paragraph() if x0 < 300 else celda_der.add_paragraph()
        
        run = p.add_run(texto_limpio)
        run.font.size = Pt(9.5)
        run.font.name = 'Arial'

    doc_word.save(docx_path)
    print(f"¡Documento Word generado de forma estructural en: {docx_path}")

if __name__ == "__main__":
    extraer_y_construir()